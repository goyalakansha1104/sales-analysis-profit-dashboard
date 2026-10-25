"""Generate an Excel sales analysis report with charts."""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList

from analysis import analyze_sales


def generate_excel_report():
    # Run sales analysis
    result = analyze_sales()

    # Output file location
    output_folder = Path("output")
    output_folder.mkdir(parents=True, exist_ok=True)
    output_file = output_folder / "sales_report.xlsx"

    # Prepare summary data
    summary = result["summary"]

    summary_df = pd.DataFrame({
        "Metric": [
            "Total Revenue",
            "Total Profit",
            "Total Orders",
            "Units Sold",
            "Average Order Value",
        ],
        "Value": [
            summary["total_revenue"],
            summary["total_profit"],
            summary["total_orders"],
            summary["units_sold"],
            summary["average_order_value"],
        ],
    })

    # Write all data to Excel sheets
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        summary_df.to_excel(
            writer, sheet_name="Executive Summary", index=False
        )
        result["raw_data"].to_excel(
            writer, sheet_name="Raw Sales Data", index=False
        )
        result["monthly"].to_excel(
            writer, sheet_name="Monthly Analysis", index=False
        )
        result["category"].to_excel(
            writer, sheet_name="Category Analysis", index=False
        )
        result["top_products"].to_excel(
            writer, sheet_name="Top Products", index=False
        )

    # Open workbook for formatting
    workbook = load_workbook(output_file)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False

        # Format column headers
        for cell in worksheet[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # Adjust column widths
        for column_cells in worksheet.columns:
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            worksheet.column_dimensions[
                column_cells[0].column_letter
            ].width = min(max_length + 3, 35)

    # Format executive summary values
    summary_sheet = workbook["Executive Summary"]

    for row in range(2, summary_sheet.max_row + 1):
        metric = summary_sheet.cell(row=row, column=1).value
        value_cell = summary_sheet.cell(row=row, column=2)

        if metric in ["Total Revenue", "Total Profit", "Average Order Value"]:
            value_cell.number_format = '#,##0.00'

        value_cell.font = Font(bold=True, size=12)

    # Format revenue and profit columns as numbers
    for sheet_name in [
        "Monthly Analysis",
        "Category Analysis",
        "Top Products",
    ]:
        worksheet = workbook[sheet_name]

        for column in range(1, worksheet.max_column + 1):
            header = worksheet.cell(row=1, column=column).value

            if header in ["revenue", "profit"]:
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(
                        row=row, column=column
                    ).number_format = '#,##0.00'

    # Format raw data date and financial columns
    raw_sheet = workbook["Raw Sales Data"]
    raw_headers = {
        cell.value: cell.column
        for cell in raw_sheet[1]
    }

    if "order_date" in raw_headers:
        date_column = raw_headers["order_date"]
        for row in range(2, raw_sheet.max_row + 1):
            raw_sheet.cell(
                row=row, column=date_column
            ).number_format = "yyyy-mm-dd"

    for header in ["revenue", "profit"]:
        if header in raw_headers:
            column = raw_headers[header]
            for row in range(2, raw_sheet.max_row + 1):
                raw_sheet.cell(
                    row=row, column=column
                ).number_format = '#,##0.00'

    # Monthly revenue and profit chart
    monthly_sheet = workbook["Monthly Analysis"]

    monthly_chart = BarChart()
    monthly_chart.type = "col"
    monthly_chart.style = 10
    monthly_chart.title = "Monthly Revenue and Profit"
    monthly_chart.y_axis.title = "Amount"
    monthly_chart.x_axis.title = "Month"
    monthly_chart.height = 8
    monthly_chart.width = 16

    monthly_data = Reference(
        monthly_sheet,
        min_col=2,
        max_col=3,
        min_row=1,
        max_row=monthly_sheet.max_row,
    )
    monthly_categories = Reference(
        monthly_sheet,
        min_col=1,
        min_row=2,
        max_row=monthly_sheet.max_row,
    )

    monthly_chart.add_data(monthly_data, titles_from_data=True)
    monthly_chart.set_categories(monthly_categories)
    monthly_sheet.add_chart(monthly_chart, "F2")

    # Category revenue chart
    category_sheet = workbook["Category Analysis"]

    category_chart = BarChart()
    category_chart.type = "bar"
    category_chart.style = 11
    category_chart.title = "Revenue by Category"
    category_chart.x_axis.title = "Revenue"
    category_chart.y_axis.title = "Category"
    category_chart.height = 8
    category_chart.width = 16

    category_data = Reference(
        category_sheet,
        min_col=2,
        min_row=1,
        max_row=category_sheet.max_row,
    )
    category_categories = Reference(
        category_sheet,
        min_col=1,
        min_row=2,
        max_row=category_sheet.max_row,
    )

    category_chart.add_data(category_data, titles_from_data=True)
    category_chart.set_categories(category_categories)
    category_sheet.add_chart(category_chart, "F2")

    # Top products revenue chart
    products_sheet = workbook["Top Products"]

    products_chart = BarChart()
    products_chart.type = "bar"
    products_chart.style = 12
    products_chart.title = "Top 10 Products by Revenue"
    products_chart.x_axis.title = "Revenue"
    products_chart.y_axis.title = "Product"
    products_chart.height = 10
    products_chart.width = 18

    products_data = Reference(
        products_sheet,
        min_col=2,
        min_row=1,
        max_row=products_sheet.max_row,
    )
    products_categories = Reference(
        products_sheet,
        min_col=1,
        min_row=2,
        max_row=products_sheet.max_row,
    )

    products_chart.add_data(products_data, titles_from_data=True)
    products_chart.set_categories(products_categories)
    products_sheet.add_chart(products_chart, "F2")

    # Save workbook
    workbook.save(output_file)

    print("\nExcel report created successfully!")
    print(f"File location: {output_file.resolve()}")
    print(f"Total Revenue: {summary['total_revenue']:.2f}")
    print(f"Total Profit: {summary['total_profit']:.2f}")
    print("Sheets: Executive Summary, Raw Sales Data, Monthly Analysis,")
    print("        Category Analysis, Top Products")


if __name__ == "__main__":
    generate_excel_report()